"""Development only: run with python-pptx and python-docx installed."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from docx import Document
root = Path(__file__).resolve().parents[1]
p = Presentation()
p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = 'Title'
s.placeholders[1].text = 'Body'
s.notes_slide.notes_text_frame.text = ''
p.save(root / 'plugins/microsoft/content/blank.pptx')
s.shapes.title.text = 'Hello World'
s.shapes.title.text_frame.paragraphs[0].runs[0].font.bold = True
s.placeholders[1].text = 'First\nSecond'
s.notes_slide.notes_text_frame.text = 'Speaker notes'
table = s.shapes.add_table(2,2,Inches(1),Inches(5),Inches(6),Inches(1)).table
for row, values in zip(table.rows, [['A','B'],['C','D']]):
    for cell, value in zip(row.cells, values): cell.text=value
p.save(root / 'test/fixtures/content/complex.pptx')
d = Document()
d.add_heading('Example', level=1)
p = d.add_paragraph()
p.add_run('Hello ').bold = True
p.add_run('World').italic = True
t = d.add_table(rows=2,cols=2)
for row, values in zip(t.rows,[['A','B'],['C','D']]):
    for cell, value in zip(row.cells,values): cell.text=value
d.sections[0].header.paragraphs[0].text = 'Preserve this header'
d.save(root / 'test/fixtures/content/complex.docx')
